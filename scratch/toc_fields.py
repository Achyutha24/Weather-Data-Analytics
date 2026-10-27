from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path
p=Path('output/Weather_Data_Storage_and_Analysis_MongoDB_Report.docx')
d=Document(p)
entries=[
 ('INTRODUCTION','1. INTRODUCTION',29),
 ('SYSTEM REQUIREMENTS AND SPECIFICATIONS','2. SYSTEM REQUIREMENTS AND SPECIFICATIONS',30),
 ('SYSTEM DESIGN','3. SYSTEM DESIGN',31),
 ('IMPLEMENTATION','4. IMPLEMENTATION',32),
 ('RESULT','5. RESULT',33),
 ('CONCLUSION','6. CONCLUSION',34),
]
# Attach stable bookmarks to the six report section headings.
for n,(name,label,ti) in enumerate(entries,1):
 target=next(x for x in d.paragraphs if x.text.strip()==label and x is not d.paragraphs[ti])
 start=OxmlElement('w:bookmarkStart'); start.set(qn('w:id'),str(2000+n)); start.set(qn('w:name'),f'ReportSection{n}')
 end=OxmlElement('w:bookmarkEnd'); end.set(qn('w:id'),str(2000+n))
 target._p.insert(1,start); target._p.append(end)
# Static TOC labels point to PAGEREF fields; Word refreshes them when the report opens.
for n,(name,label,ti) in enumerate(entries,1):
 para=d.paragraphs[ti]
 para.clear()
 para.add_run(label+'  ')
 dots=para.add_run('.'*55+' ')
 fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),f'PAGEREF ReportSection{n} \\h')
 rr=OxmlElement('w:r'); t=OxmlElement('w:t'); t.text='1'; rr.append(t); fld.append(rr); para._p.append(fld)
settings=d.settings.element
u=settings.find(qn('w:updateFields'))
if u is None: u=OxmlElement('w:updateFields'); settings.append(u)
u.set(qn('w:val'),'true')
d.save(p)
print('dynamic TOC fields added',len(entries))
