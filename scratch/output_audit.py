from docx import Document
from docx.oxml.ns import qn
from pathlib import Path
from zipfile import ZipFile
import hashlib,re
p=Path(r'C:\Users\achyu\Desktop\BDA\output\Weather_Data_Storage_and_Analysis_MongoDB_Report.docx')
d=Document(p)
print('bytes',p.stat().st_size,'sha256',hashlib.sha256(p.read_bytes()).hexdigest(),'paras',len(d.paragraphs),'tables',len(d.tables),'sections',len(d.sections))
for i,x in enumerate(d.paragraphs):
 if x.text.strip(): print(f'{i:03d} [{x.style.name}] {x.text[:175]}')
print('TABLES')
for i,t in enumerate(d.tables):
 print('T',i,'rows',len(t.rows),'cols',len(t.columns),'sample',repr(' | '.join(c.text.replace('\n',' / ') for c in t.rows[0].cells)))
 print('  ',repr(' | '.join(c.text.replace('\n',' / ') for r in t.rows[1:3] for c in r.cells)[:220]))
body='\n'.join([x.text for x in d.paragraphs]+[c.text for t in d.tables for r in t.rows for c in r.cells])
for s in d.sections:
 body+='\n'+'\n'.join(p.text for p in s.header.paragraphs+s.footer.paragraphs)
print('OBSOLETE', {x:len(re.findall(x,body,re.I)) for x in ['VirtualBox','Ubuntu','virtual machine','Guest Additions','Virtual Machine Installation']})
with ZipFile(p) as z:
 print('zip test',z.testzip(),'parts',len(z.namelist()))
 xml=z.read('word/document.xml'); print('drawing',xml.count(b'<w:drawing'),'pict',xml.count(b'<w:pict'))
 print('admin excerpts',[x for x in [q.text for q in d.paragraphs] if any(k in x for k in ['ACHYUTHA','1VE23CA002','Vimala Devi','2025 - 2026','Sri Venkateshwara'])])
