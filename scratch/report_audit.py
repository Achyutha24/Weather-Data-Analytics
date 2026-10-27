import docx, zipfile, re
p='output/Weather_Data_Storage_and_Analysis_MongoDB_Report.docx'
d=docx.Document(p)
for i,x in enumerate(d.paragraphs):
 t=x.text.strip()
 if t:
  print(i, 'DRAW' if 'drawing' in x._p.xml else '', repr(t[:180]))
print('SECTION BREAKS',sum(1 for x in d.paragraphs if 'sectPr' in x._p.xml))
print('CSV lines',sum(1 for _ in open('data/sample_weather.csv',encoding='utf-8-sig')))
