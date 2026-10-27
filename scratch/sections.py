from docx import Document
from docx.oxml.ns import qn
from lxml import etree
for path in ['C:/Users/achyu/Downloads/cc report task 1.docx','output/Weather_Data_Storage_and_Analysis_MongoDB_Report.docx']:
 d=Document(path); print('\nDOC',path,'paragraphs',len(d.paragraphs),'sections',len(d.sections))
 for i,p in enumerate(d.paragraphs):
  sect=p._p.pPr.sectPr if p._p.pPr is not None else None
  if sect is not None:
   print('SECT at para',i,repr(p.text[:80]),'type',sect.find(qn('w:type')).get(qn('w:val')) if sect.find(qn('w:type')) is not None else 'nextPage','border',etree.tostring(sect.find(qn('w:pgBorders')))[:300] if sect.find(qn('w:pgBorders')) is not None else None)
 print('final sect borders:',[etree.tostring(s._sectPr.find(qn('w:pgBorders')))[:300] if s._sectPr.find(qn('w:pgBorders')) is not None else None for s in d.sections])
