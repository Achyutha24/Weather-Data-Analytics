from docx import Document
from collections import Counter
from lxml import etree
for f in [r'C:\Users\achyu\Downloads\cc report task 1.docx',r'output\Weather_Data_Storage_and_Analysis_MongoDB_Report.docx']:
 d=Document(f); print('\n',f)
 for i in [0,5,21,41,85,124,148,237,343]:
  x=d.paragraphs[i]._p; print(i,'tags',Counter(etree.QName(n).localname for n in x.iter()),'runs',len(d.paragraphs[i].runs),'drawing',len(x.findall('.//{*}drawing')),'pict',len(x.findall('.//{*}pict')))
  if i==5: print(etree.tostring(x,encoding='unicode')[-700:])
