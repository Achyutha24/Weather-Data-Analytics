from docx import Document
from docx.oxml.ns import qn
from pathlib import Path
import zipfile, hashlib

src = Path(r'C:\Users\achyu\Downloads\cc report task 1.docx')
d = Document(src)
print('SOURCE', src, 'bytes', src.stat().st_size, 'sha256', hashlib.sha256(src.read_bytes()).hexdigest())
print('paragraphs',len(d.paragraphs),'tables',len(d.tables),'sections',len(d.sections))
for i,p in enumerate(d.paragraphs):
    if p.text.strip():
        ppr=p._p.pPr
        num=ppr.find(qn('w:numPr')) if ppr is not None else None
        numbering = ' num' if num is not None else ''
        runs=[r for r in p.runs if r.text.strip()]
        rf=[]
        for r in runs[:2]:
            rf.append(f'{r.font.name}:{r.font.size.pt if r.font.size else None}pt:b{r.bold}')
        print(f'{i:03d} | {p.style.name}{numbering} | {p.alignment} | {";".join(rf)} | {p.text[:230]!r}')
print('TABLES')
for ti,t in enumerate(d.tables):
    print('T',ti,len(t.rows),len(t.columns),'style',t.style.name if t.style else None,'text',repr(' | '.join(c.text.replace('\n',' / ') for r in t.rows[:2] for c in r.cells)[:240]))
print('SECTIONS')
for si,s in enumerate(d.sections):
    print(si,'size',round(s.page_width.inches,2),round(s.page_height.inches,2),'margins',round(s.top_margin.inches,2),round(s.bottom_margin.inches,2),round(s.left_margin.inches,2),round(s.right_margin.inches,2),'header/footer',round(s.header_distance.inches,2),round(s.footer_distance.inches,2),'different_first',s.different_first_page_header_footer)
    print('  header', [p.text for p in s.header.paragraphs], 'footer', [p.text for p in s.footer.paragraphs], 'first', [p.text for p in s.first_page_header.paragraphs], [p.text for p in s.first_page_footer.paragraphs])
with zipfile.ZipFile(src) as z:
    print('PACKAGE')
    for n in z.namelist():
        if n in {'word/document.xml','word/styles.xml','word/numbering.xml','word/settings.xml'} or n.startswith('word/header') or n.startswith('word/footer'):
            print(n,z.getinfo(n).file_size)
    xml=z.read('word/document.xml')
    print('drawing count',xml.count(b'<w:drawing'),'pict count',xml.count(b'<w:pict'),'section count',xml.count(b'<w:sectPr'))
