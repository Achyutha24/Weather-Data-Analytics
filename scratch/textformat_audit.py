from docx import Document
from docx.oxml.ns import qn
from pathlib import Path
for path in [Path(r'C:\Users\achyu\Downloads\cc report task 1.docx'),Path('output/Weather_Data_Storage_and_Analysis_MongoDB_Report.docx')]:
 d=Document(path); print('\n',path, path.stat().st_size, path.stat().st_mtime)
 print('cover',repr(d.paragraphs[11].text))
 for i in list(range(21,25))+list(range(29,35))+[39,85,124,148,237,343]:
  p=d.paragraphs[i]; print(i,repr(p.text),p.style.name,'num',p._p.pPr.numPr.xml if p._p.pPr is not None and p._p.pPr.numPr is not None else None,'runs',[(r.text,r.bold,r.font.size.pt if r.font.size else None) for r in p.runs])
  if i in [29,30]: print(p._p.xml[-900:])
