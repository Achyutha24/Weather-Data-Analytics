from docx import Document
import zipfile,os,re
from lxml import etree
p='output/Weather_Data_Storage_and_Analysis_MongoDB_Report.docx'
d=Document(p)
print('bytes',os.path.getsize(p),'paras',len(d.paragraphs),'tables',len(d.tables),'sections',len(d.sections))
print('table dimensions',[(len(t.rows),len(t.columns)) for t in d.tables])
texts=[x.text for x in d.paragraphs]
alltext='\n'.join(texts+[p.text for s in d.sections for a in [s.header,s.footer,s.first_page_header,s.first_page_footer] for p in a.paragraphs])
for term in ['VirtualBox','Ubuntu','virtual machine','Guest Additions','Virtual Machine Installation']:
 print('obsolete',term,len(re.findall(term,alltext,re.I)))
print('admin',[(i,texts[i]) for i in [11,13,14,15,16]])
print('headings',[x.text for x in d.paragraphs if x.text.strip() and (x.style.name.startswith('Heading') or re.match(r'^[1-6]\. ',x.text))])
print('text extracted chars',len(alltext))
z=zipfile.ZipFile(p); print('zip test',z.testzip())
ns={'r':'http://schemas.openxmlformats.org/package/2006/relationships'}
rel=etree.fromstring(z.read('word/_rels/document.xml.rels'))
print('image rels',[x.get('Target') for x in rel if x.get('Type','').endswith('/image')])
print('drawings',z.read('word/document.xml').count(b'<w:drawing'))
